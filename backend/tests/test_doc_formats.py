"""Library documents beyond PDF: Word, HTML, Markdown and text become cited, chunked documents."""
import io
import os
import tempfile
import zipfile

tmp = tempfile.mkdtemp()
os.environ.update({"APP_ENV": "test", "DATABASE_URL": f"sqlite:///{tmp}/t.db", "STORAGE_DIR": tmp, "AI_PROVIDER": "mock"})

import docx  # noqa: E402
import pytest  # noqa: E402

from app.core import doc_formats as df  # noqa: E402
from app.core.chunking import chunk_pages  # noqa: E402

BODY = "Little's Law relates the average work in process to throughput and cycle time in any stable system. " * 3


def word_file() -> bytes:
    d = docx.Document()
    d.core_properties.title = "Factory Physics Notes"
    d.add_heading("Flow basics", level=1)
    d.add_paragraph(BODY)
    d.add_heading("Little's Law", level=2)
    d.add_paragraph("Work in process equals throughput times cycle time. " * 4)
    t = d.add_table(rows=2, cols=3)
    for r, row in enumerate([["Station", "Rate", "Utilisation"], ["Drill", "12", "0.85"]]):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    d.add_paragraph("Bullet about variability and queues in production systems", style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


@pytest.mark.parametrize("data, name, fmt", [
    (b"%PDF-1.7 ...", "x.pdf", "pdf"),
    (b"<!doctype html><html><body><p>Hi</p></body></html>", "page.bin", "html"),
    (b"# Title\n\nSome text here.", "notes.txt", "markdown"),
    (b"Plain notes about queues.", "notes.txt", "text"),
    ("Café notes, naïve résumé".encode("latin-1"), "n.txt", "text"),
    (b"\x89PNG\r\n\x1a\n\x00\x00", "image.png", None),
    (b"\x00\x01\x02binary", "file.txt", None),
])
def test_detect_by_content(data, name, fmt):
    assert df.detect(data, name) == fmt


def test_detect_word_and_reject_other_zips():
    assert df.detect(word_file(), "anything.bin") == "docx"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("hello.txt", "x")
    assert df.detect(buf.getvalue(), "x.docx") is None


def test_word_keeps_headings_tables_and_lists_in_order():
    r = df.extract_text_document(word_file(), "docx")
    assert r.title == "Factory Physics Notes"
    text = "\n".join(p.text for p in r.pages)
    assert text.index("# Flow basics") < text.index("## Little's Law") < text.index("| Station | Rate | Utilisation |")
    assert "| Drill | 12 | 0.85 |" in text and "• Bullet about variability" in text
    chunks = chunk_pages(r.pages)
    assert chunks[0].section_label == "Flow basics"
    assert any(c.section_label == "Flow basics > Little's Law" and "| Drill | 12 | 0.85 |" in c.text for c in chunks)


def test_html_headings_and_tables():
    html = (b"<html><head><title>Ops notes</title><script>x()</script></head><body><nav>menu</nav>"
            b"<h1>Queues</h1><p>" + BODY.encode() + b"</p><h2>Data</h2><table><tr><th>Station</th><th>Rate</th></tr>"
            b"<tr><td>Drill</td><td>12</td></tr></table></body></html>")
    r = df.extract_text_document(html, "html")
    text = r.pages[0].text
    assert r.title == "Ops notes" and "x()" not in text and "menu" not in text
    assert "# Queues" in text and "## Data" in text and "| Drill | 12 |" in text


def test_markdown_tables_links_and_long_files_split_into_pages():
    md = "# Notes\n\n" + "\n\n".join(f"Paragraph {i}: {BODY}" for i in range(20)) + \
         "\n\n| A | B |\n|---|---|\n| 1 | 2 |\n\nSee [the guide](http://x.example) for **more**."
    r = df.extract_text_document(md.encode(), "markdown")
    assert r.page_count > 1 and all(len(p.text) <= df.PAGE_CHARS + 400 for p in r.pages)
    last = r.pages[-1].text
    assert "| 1 | 2 |" in last and "---" not in last and "See the guide for more." in last
    assert r.title == "Notes"


def test_empty_and_oversized_files_are_refused():
    with pytest.raises(df.PdfExtractionError):
        df.extract_text_document(b"   \n  ", "text")
    with pytest.raises(df.PdfExtractionError):
        df.extract_text_document(("x " * 2000 + "\n\n").encode() * 50, "text", max_pages=3)


def test_upload_word_then_ask_with_citations(client):
    r = client.post("/api/auth/signup", json={"email": "w@example.com", "password": "correct-horse-1"})
    h = {"Authorization": f"Bearer {r.json()['token']}"}
    up = client.post("/api/documents?rights_confirmed=true", headers=h,
                     files={"file": ("notes.docx", word_file(), "application/octet-stream")})
    assert up.status_code == 201
    doc = client.get(f"/api/documents/{up.json()['id']}", headers=h).json()
    assert doc["status"] == "ready" and doc["page_count"] >= 1
    ans = client.post("/api/tutor/ask", headers=h, json={"question": "What does Little's Law relate?", "document_ids": [doc["id"]]}).json()
    assert ans["status"] == "verified" and ans["citations"]
    f = client.get(f"/api/documents/{doc['id']}/file", headers=h)
    assert f.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument")
    t = client.post("/api/documents?rights_confirmed=true", headers=h,
                    files={"file": ("notes.md", ("# Queues\n\n" + BODY).encode(), "text/markdown")})
    assert t.status_code == 201
    assert client.get(f"/api/documents/{t.json()['id']}", headers=h).json()["status"] == "ready"
