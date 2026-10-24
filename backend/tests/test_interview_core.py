"""Pure interview-prep logic: validation, gap rules, priority, engine, deep dives (PRD §7)."""
import io
from datetime import datetime, timedelta, timezone

import pytest

from app.core.interview_deepdive import analyse_claim, merge_questions, template_questions
from app.core.interview_engine import (
    Candidate,
    ResumeBulletRef,
    SessionState,
    TurnState,
    band_verdict,
    build_summary,
    needs_follow_up,
    next_step,
    normalize_evaluation,
    rubric_for,
    select_question,
)
from app.core.interview_matrix import (
    UNASSESSED,
    AssessmentSignal,
    InterviewSignal,
    band_priority,
    build_skill,
    concept_matches,
    gap_type,
    priority_factors,
    priority_score,
    resume_strength,
    skill_listed,
    urgency,
    validate_judgements,
)
from app.core.interview_parse import (
    importance_from_text,
    normalize_job_edit,
    normalize_resume_edit,
    validate_job,
    validate_resume,
)
from app.core.interview_plan import TurnScore, build_plan, build_readiness
from app.core.interview_text import InputError, UnsupportedFileType, extract_resume_file, normalize_source, prepare_text
from app.core.mastery import AttemptEvidence
from tests.test_interview_samples import SAMPLE_JD, SAMPLE_RESUME

NOW = datetime(2026, 10, 10, tzinfo=timezone.utc)


# ---------------- input extraction ----------------

def test_normalize_source_unifies_bullets_and_keeps_lines():
    out = normalize_source("Experience\r\n\r\n\r\n\r\n●  Built   things\n▪ Led team")
    assert out == "Experience\n\n• Built things\n• Led team"


def test_prepare_text_rejects_tiny_input_with_actionable_message():
    with pytest.raises(InputError, match="Paste the full text"):
        prepare_text("Python, SQL", "resume")


def test_prepare_text_caps_long_input_with_warning():
    out = prepare_text("Did a thing well.\n" * 3000, "job description")
    assert len(out.text) <= 20_000 and out.warnings


def test_unsupported_file_type():
    with pytest.raises(UnsupportedFileType, match="PDF or DOCX"):
        extract_resume_file(b"hello world", "resume.txt")


def test_extract_docx_marks_list_paragraphs_as_bullets():
    import docx

    d = docx.Document()
    d.add_paragraph("Priya Sharma")
    d.add_paragraph("Experience")
    d.add_paragraph("Improved demand forecast accuracy by 8% using gradient boosting", style="List Bullet")
    d.add_paragraph("Built SQL pipelines in BigQuery for weekly sales dashboards", style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    out = extract_resume_file(buf.getvalue(), "cv.docx")
    assert "• Improved demand forecast accuracy by 8%" in out.text


def test_extract_pdf_resume():
    from tests.conftest import make_pdf

    pdf = make_pdf([SAMPLE_RESUME.splitlines()])
    out = extract_resume_file(pdf, "cv.pdf")
    assert "Northwind Retail" in out.text
    # reportlab's bullet glyph extracts as a control character; it must become a real bullet
    assert "\n• Improved demand forecast accuracy by 8%" in out.text
    assert "\x7f" not in out.text


def test_scanned_pdf_is_rejected_with_hint():
    from tests.conftest import make_pdf

    with pytest.raises(InputError, match="scanned"):
        extract_resume_file(make_pdf([["x"]]), "scan.pdf")


# ---------------- parse validation ----------------

def _raw_resume(extra_bullets=()):
    return {
        "name": "Priya Sharma",
        "summary": None,
        "roles": [{
            "title": "Senior Data Analyst", "company": "Northwind Retail", "start": "Mar 2022", "end": "Present",
            "bullets": [
                "• Improved demand forecast accuracy by 8% by replacing a moving-average model with gradient boosting in Python",
                "Built SQL pipelines in BigQuery that feed weekly sales dashboards used by 40 store managers",
                *extra_bullets,
            ],
        }],
        "skills": ["Python", "SQL", "Rust"],
        "education": ["BSc Mathematics, University of Leeds, 2019"],
        "certifications": [],
    }


def test_validation_drops_hallucinated_bullets_and_assigns_ids():
    out = validate_resume(_raw_resume(["Grew revenue by 300% across 12 international markets as head of sales"]), SAMPLE_RESUME)
    bullets = out.parsed["roles"][0]["bullets"]
    assert [b["id"] for b in bullets] == ["r1b1", "r1b2"]
    assert bullets[0]["text"].startswith("Improved demand forecast")  # marker stripped
    assert any("not in your resume" in w for w in out.warnings)
    assert out.parsed["skills"] == ["Python", "SQL"]  # "Rust" was never in the resume
    assert out.parsed["education"] == [{"id": "e1", "text": "BSc Mathematics, University of Leeds, 2019"}]


def test_validation_tolerates_pdf_whitespace_artifacts():
    raw = _raw_resume()
    raw["roles"][0]["bullets"] = ["Improved demand fore-\ncast accuracy by 8% by replacing a moving-average model with gradient boosting in  Python"]
    out = validate_resume(raw, SAMPLE_RESUME)
    assert len(out.parsed["roles"][0]["bullets"]) == 1


def test_validation_rejects_an_embellished_bullet():
    raw = _raw_resume()
    raw["roles"][0]["bullets"] = ["Improved demand forecast accuracy by 18% by replacing a legacy model with deep learning in Python"]
    out = validate_resume(raw, SAMPLE_RESUME)
    assert out.parsed["roles"][0]["bullets"] == []  # the role header is real; the inflated bullet is not
    assert any("not in your resume" in w for w in out.warnings)


def test_validation_rejects_non_object():
    with pytest.raises(ValueError):
        validate_resume(["not", "an", "object"], SAMPLE_RESUME)


def test_job_validation_drops_unquoted_competencies_and_coerces_enums():
    raw = {
        "role_title": "Analytics Manager", "company": "Fabrikam Foods",
        "competencies": [
            {"name": "SQL", "category": "technical", "importance": "high", "required": True, "evidence_quote": "5+ years of experience with SQL and Python"},
            {"name": "Rust", "category": "technical", "importance": "critical", "evidence_quote": "Expert Rust programming required"},
            {"name": "Kubernetes", "category": "weird", "importance": "???", "evidence_quote": "Experience with Kubernetes is a plus"},
        ],
        "responsibilities": ["Build demand forecasting models that drive inventory decisions", "Invented duty"],
    }
    out = validate_job(raw, SAMPLE_JD)
    comps = out.parsed["competencies"]
    assert [c["name"] for c in comps] == ["SQL", "Kubernetes"]
    assert [c["id"] for c in comps] == ["c1", "c2"]
    assert comps[1]["importance"] == "low" and comps[1]["category"] == "tool" and comps[1]["required"] is False
    assert out.parsed["responsibilities"] == ["Build demand forecasting models that drive inventory decisions"]
    assert any("Rust" in w for w in out.warnings)


@pytest.mark.parametrize("text, expected", [
    ("Must have experience with A/B testing", "critical"),
    ("SQL is required", "critical"),
    ("Experience with Tableau preferred", "medium"),
    ("Kubernetes is nice to have", "low"),
    ("Experience with Docker is a plus", "low"),
    ("Experience with Python", "high"),
])
def test_importance_from_wording(text, expected):
    assert importance_from_text(text, "high") == expected


def test_resume_edit_keeps_ids_and_mints_missing_ones():
    edited = normalize_resume_edit({
        "roles": [{"id": "r1", "title": "Analyst", "bullets": [{"id": "r1b1", "text": "Kept"}, {"id": "", "text": "New bullet"}, {"id": "r1b1", "text": "Dup id"}]}],
        "skills": ["SQL", "SQL", " "], "education": [], "certifications": [],
    })
    ids = [b["id"] for b in edited["roles"][0]["bullets"]]
    assert ids[0] == "r1b1" and len(set(ids)) == 3
    assert edited["skills"] == ["SQL"]


def test_edits_replace_client_ids_without_stealing_existing_ones():
    # The web client sends "role-<uuid>", "b-<uuid>", "comp-<uuid>" for rows the user adds, and may list them first.
    resume = normalize_resume_edit({
        "roles": [{"id": "role-1f2e", "title": "New", "bullets": [{"id": "b-9a", "text": "Added"}]},
                  {"id": "r1", "title": "Old", "bullets": [{"id": "r1b1", "text": "Kept"}]}],
        "education": [{"id": "edu-x", "text": "BSc"}],
    })
    assert [r["id"] for r in resume["roles"]][1] == "r1"
    assert resume["roles"][1]["bullets"][0]["id"] == "r1b1"
    ids = [r["id"] for r in resume["roles"]] + [b["id"] for r in resume["roles"] for b in r["bullets"]] + [e["id"] for e in resume["education"]]
    assert len(set(ids)) == len(ids) and not any("-" in i for i in ids)
    job = normalize_job_edit({"role_title": "X", "competencies": [
        {"id": "comp-123e4567-e89b-12d3-a456-426614174000", "name": "New"}, {"id": "c1", "name": "Old"}, {"id": "", "name": "Blank"}]})
    cids = [c["id"] for c in job["competencies"]]
    assert cids[1] == "c1" and len(set(cids)) == 3 and not any("-" in i for i in cids)


# ---------------- matching and gaps ----------------

def test_judgements_drop_unknown_ids():
    raw = {"matches": [
        {"competency_id": "c1", "evidence": [{"bullet_id": "r1b1", "strength": "strong"}, {"bullet_id": "r9b9", "strength": "strong"},
                                            {"bullet_id": "r1b2", "strength": "certain"}]},
        {"competency_id": "c99", "evidence": [{"bullet_id": "r1b1", "strength": "strong"}]},
    ]}
    assert validate_judgements(raw, {"c1"}, {"r1b1", "r1b2"}) == {"c1": [("r1b1", "strong")]}
    with pytest.raises(ValueError):
        validate_judgements({"oops": []}, {"c1"}, set())


def test_skills_list_only_is_partial_never_strong():
    assert skill_listed("SQL", ["Python", "SQL"])
    assert resume_strength([], listed=True) == "partial"
    comp = {"id": "c1", "name": "SQL", "category": "technical", "importance": "critical"}
    strong_quiz = AssessmentSignal("strong", 0.9, 5, "Strong.")
    for assessed in (UNASSESSED, strong_quiz):
        skill = build_skill(comp, [], {}, ["SQL"], assessed, None, None)
        assert skill.resume_strength == "partial"
        assert skill.gap_type != "strong"


def test_no_assessment_is_unverified_never_weak_or_knowledge_gap():
    for strength in ("strong", "partial", "none"):
        for importance in ("critical", "low"):
            g = gap_type(strength, UNASSESSED, None, importance)
            assert g not in ("knowledge_gap", "application_gap")
    assert gap_type("strong", UNASSESSED, None, "high") == "unverified"
    comp = {"id": "c1", "name": "SQL", "category": "technical", "importance": "high"}
    assert build_skill(comp, [], {}, [], UNASSESSED, None, None).assessed_level == "unverified"


@pytest.mark.parametrize("strength, assessed, interview, expected", [
    ("strong", AssessmentSignal("weak", 0.3, 4, ""), None, "knowledge_gap"),
    ("none", AssessmentSignal("developing", 0.6, 4, ""), None, "application_gap"),
    ("strong", AssessmentSignal("strong", 0.9, 4, "", applied_count=2, applied_mean=0.4), None, "application_gap"),
    ("strong", AssessmentSignal("strong", 0.9, 4, ""), None, "strong"),
    ("partial", AssessmentSignal("strong", 0.9, 4, ""), None, "evidence_gap"),
    ("strong", UNASSESSED, InterviewSignal(2, 0.4), "communication_gap"),
    ("strong", UNASSESSED, InterviewSignal(2, 0.85), "strong"),
    ("partial", UNASSESSED, None, "evidence_gap"),
    ("none", UNASSESSED, None, "evidence_gap"),
])
def test_gap_categories(strength, assessed, interview, expected):
    assert gap_type(strength, assessed, interview, "high") == expected


def test_low_importance_without_evidence_is_stretch():
    assert gap_type("none", UNASSESSED, None, "low") == "stretch"


def test_concept_matching_uses_content_tokens():
    assert concept_matches("SQL", "SQL")
    assert concept_matches("Demand forecasting", "Forecasting demand")
    assert not concept_matches("Stakeholder management", "Project management")
    assert not concept_matches("Python", "Little's Law")


# ---------------- priority ----------------

def test_priority_factors_are_unit_interval():
    for gap in ("knowledge_gap", "strong", "stretch", "unverified"):
        for days in (None, -3, 0, 5, 90):
            f = priority_factors(gap, "none", "critical", "low", days)
            assert set(f) == {"importance", "gap_severity", "evidence_confidence", "urgency", "feasibility"}
            assert all(0.0 <= v <= 1.0 for v in f.values())


def test_priority_monotone_in_importance_and_urgency():
    def score(importance="high", days=10, gap="evidence_gap"):
        return priority_score(priority_factors(gap, "none", importance, "medium", days))

    order = ["low", "medium", "high", "critical"]
    scores = [score(importance=i) for i in order]
    assert scores == sorted(scores)
    bands = [band_priority(s) for s in scores]
    rank = {"low": 0, "medium": 1, "high": 2, "critical": 3}
    assert [rank[b] for b in bands] == sorted(rank[b] for b in bands)
    assert score(days=2) > score(days=20) > score(days=60)
    for gap in ("knowledge_gap", "application_gap", "evidence_gap", "unverified", "strong"):
        assert score(gap="knowledge_gap") >= score(gap=gap)  # severity orders the same way
    assert urgency(0) == 1.0 and urgency(None) == 0.5 and urgency(100) == 0.1


def test_priority_bands_for_typical_cases():
    def prio(gap, importance, conf="low", strength="none"):
        return band_priority(priority_score(priority_factors(gap, strength, importance, conf, None)))

    assert prio("evidence_gap", "critical") == "critical"
    assert prio("knowledge_gap", "medium", "medium") == "high"
    assert prio("strong", "critical", "high", "strong") == "low"
    assert prio("stretch", "low") == "low"


def test_big_gap_close_to_interview_is_less_feasible():
    far = priority_factors("stretch", "none", "high", "low", 30)["feasibility"]
    near = priority_factors("stretch", "none", "high", "low", 1)["feasibility"]
    assert near < far


def test_user_override_drives_order_and_plan_but_keeps_computed_priority():
    comp = {"id": "c1", "name": "Kubernetes", "category": "tool", "importance": "low"}
    s = build_skill(comp, [], {}, [], UNASSESSED, None, None, override={"priority": "critical"})
    assert s.priority == "low" and s.effective_priority == "critical"
    plan = build_plan([s])
    assert plan[0]["competency_id"] == "c1"


# ---------------- plan and readiness ----------------

def _skill(cid, name, category, gap, importance="high", strength="none", assessed="unverified"):
    comp = {"id": cid, "name": name, "category": category, "importance": importance}
    signal = UNASSESSED if assessed == "unverified" else AssessmentSignal(assessed, 0.5, 4, "")
    s = build_skill(comp, [], {}, [], signal, None, None)
    s.gap_type, s.resume_strength = gap, strength
    return s


def test_plan_maps_gap_types_to_activities_and_ends_with_mock_interview():
    plan = build_plan([
        _skill("c1", "SQL", "technical", "knowledge_gap"),
        _skill("c2", "Python", "tool", "application_gap"),
        _skill("c3", "Stakeholder management", "behavioral", "evidence_gap"),
    ])
    by_comp = {}
    for item in plan:
        by_comp.setdefault(item["competency_id"], []).append(item["activity"])
    assert by_comp["c1"] == ["learn", "practice"]
    assert by_comp["c2"] == ["drill"]
    assert by_comp["c3"] == ["story"]
    assert plan[-1]["activity"] == "mock_interview" and plan[-1]["competency_id"] is None
    assert len({i["id"] for i in plan}) == len(plan)


def test_readiness_has_no_overall_number_and_needs_evidence():
    skills = [_skill("c1", "SQL", "technical", "unverified", strength="strong")]
    dims = build_readiness(skills, [])
    assert [d["dimension"] for d in dims] == ["role_knowledge", "applied_problem_solving", "resume_evidence", "behavioral_examples", "communication"]
    assert all(set(d) == {"dimension", "level", "evidence_count", "explanation"} for d in dims)
    assert dims[0]["level"] == "unverified" and dims[2]["level"] == "strong"
    dims = build_readiness(skills, [TurnScore("behavioral", 0.9), TurnScore("behavioral", 0.8), TurnScore("technical", 0.3)])
    behavioral = next(d for d in dims if d["dimension"] == "behavioral_examples")
    assert behavioral["level"] == "strong" and behavioral["evidence_count"] == 2


# ---------------- deep dive ----------------

def test_deep_dive_metric_templates_cover_definition_baseline_method_attribution():
    qs = template_questions("Improved forecast accuracy by 8% by switching to gradient boosting")
    text = " ".join(qs).lower()
    assert '"8%"' in qs[0]
    for probe in ("measured", "baseline", "period", "method", "other factors"):
        assert probe in text


def test_deep_dive_detects_from_to_metrics_and_skips_plain_bullets():
    assert analyse_claim("Reduced reporting time from 3 days to 4 hours").metric == "from 3 days to 4 hours"
    assert template_questions("Responsible for the weekly report") == []
    assert any("personal role" in q for q in template_questions("Led a cross-functional pricing review"))


def test_merge_questions_dedupes_and_caps():
    merged = merge_questions(["A question here?"], ["a question here", "Another one?"] + [f"Q{i}?" for i in range(10)])
    assert merged[:2] == ["A question here?", "Another one?"] and len(merged) == 6


# ---------------- interview engine ----------------

CANDIDATES = [
    Candidate("c1", "SQL", "technical", "high", "unverified", ("Built SQL pipelines",)),
    Candidate("c2", "Stakeholder management", "behavioral", "critical", "evidence_gap"),
    Candidate("c3", "Supply chain", "domain", "medium", "stretch"),
]
BULLETS = [
    ResumeBulletRef("r1b1", "Improved demand forecast accuracy by 8% with gradient boosting", "Analyst", "c1", "SQL"),
    ResumeBulletRef("r1b2", "Responsible for weekly reporting", "Analyst"),
]


def _answered(kind, category, score, missing=(), cid="c1"):
    ev = {"verdict": band_verdict(score), "score": score, "missing_details": list(missing), "criteria": [], "improvements": []}
    return TurnState(kind, category, f"Q {kind} {category} {score}", cid, "SQL", "answer", ev)


def test_follow_up_only_when_needed_and_at_most_once():
    s = SessionState("mixed", 5, [_answered("question", "technical", 0.9)])
    assert not needs_follow_up(s)
    s = SessionState("mixed", 5, [_answered("question", "technical", 0.9, missing=["The baseline"])])
    assert next_step(s, CANDIDATES, BULLETS).action == "follow_up"
    s = SessionState("mixed", 5, [_answered("question", "technical", 0.3)])
    assert next_step(s, CANDIDATES, BULLETS).action == "follow_up"
    s.turns.append(_answered("follow_up", "technical", 0.2, missing=["Still missing"]))
    assert next_step(s, CANDIDATES, BULLETS).action == "question"


def test_unassessed_answer_gets_no_follow_up():
    t = TurnState("question", "technical", "Q", "c1", "SQL", "ok", {"verdict": "unassessed", "score": 0.0, "missing_details": []})
    assert not needs_follow_up(SessionState("mixed", 3, [t]))


def test_session_completes_at_target_and_requires_an_answer():
    s = SessionState("mixed", 1, [_answered("question", "technical", 0.9)])
    assert next_step(s, CANDIDATES, BULLETS).action == "complete"
    open_turn = TurnState("question", "technical", "Q", "c1", "SQL")
    with pytest.raises(ValueError):
        next_step(SessionState("mixed", 3, [open_turn]), CANDIDATES, BULLETS)


def test_question_selection_balances_categories_and_competencies():
    s = SessionState("mixed", 5)
    seen = []
    for _ in range(5):
        plan = select_question(s, CANDIDATES, BULLETS)
        seen.append((plan.category, plan.competency_id))
        s.turns.append(TurnState("question", plan.category, f"question {len(seen)} {plan.bullet}", plan.competency_id, plan.competency_name, "a", {"score": 1.0}))
    cats = [c for c, _ in seen]
    assert cats == ["technical", "behavioral", "resume_deep_dive", "role_specific", "case"]
    assert seen[1][1] == "c2"  # behavioral picks the behavioral competency
    assert {cid for _, cid in seen} >= {"c1", "c2", "c3"}


def test_resume_focus_falls_back_without_probeable_bullets():
    plan = select_question(SessionState("resume", 3), CANDIDATES, [ResumeBulletRef("r1b1", "Responsible for reports", "A")])
    assert plan.category == "behavioral"
    plan = select_question(SessionState("resume", 3), CANDIDATES, BULLETS)
    assert plan.category == "resume_deep_dive" and plan.bullet == BULLETS[0].text


def test_score_is_recomputed_from_criteria_not_model_claim():
    rubric = rubric_for("behavioral")
    raw = {
        "score": 1.0, "verdict": "strong",
        "criteria": [{"criterion": r["criterion"], "status": "met" if i < 2 else "missing", "evidence": ""} for i, r in enumerate(rubric)],
        "strengths": ["Clear context"], "improvements": ["Say you saved $2M"], "missing_details": ["The result you achieved"],
        "feedback": "ok",
    }
    ev = normalize_evaluation(raw, "behavioral", "Tell me about a time", "We had a problem and my task was to fix it.")
    total = sum(r["points"] for r in rubric)
    assert ev["score"] == round(2 / total, 3)
    assert ev["verdict"] == "needs_work"
    assert ev["star"] == {"situation": True, "task": True, "action": False, "result": False, "reflection": False}
    assert ev["improvements"] == []  # cited a figure the candidate never gave
    assert ev["missing_details"] == ["The result you achieved"]


def test_unparseable_grade_is_unassessed():
    ev = normalize_evaluation({"feedback": "??"}, "technical", "Q", "A")
    assert ev["verdict"] == "unassessed" and ev["score"] == 0.0 and ev["star"] is None
    assert len(ev["criteria"]) == len(rubric_for("technical"))


@pytest.mark.parametrize("score, verdict", [(0.85, "strong"), (0.8, "strong"), (0.65, "adequate"), (0.59, "needs_work")])
def test_verdict_bands(score, verdict):
    assert band_verdict(score) == verdict


def test_summary_from_evaluated_turns_only():
    turns = [_answered("question", "technical", 0.4), _answered("follow_up", "technical", 0.8),
             _answered("question", "behavioral", 0.9, cid="c2"),
             TurnState("question", "case", "Q", "c3", "Supply chain", "x", {"verdict": "unassessed", "score": 0.0})]
    s = build_summary(turns)
    assert [r["competency_id"] for r in s["by_competency"]] == ["c1", "c2"]
    assert s["by_competency"][0] == {"competency_id": "c1", "name": "SQL", "score": 0.6, "attempts": 2}
    assert "3 evaluated" in s["overall"]
    assert build_summary([])["by_competency"] == []


def test_assessment_signal_pools_matching_quiz_concepts():
    from app.core.interview_matrix import assessment_signal

    old = NOW - timedelta(days=1)
    attempts = {"SQL joins": [AttemptEvidence(0.2, "mcq", "easy", old)] * 3, "Little's Law": [AttemptEvidence(1.0, "mcq", "easy", old)] * 3}
    sig = assessment_signal("SQL", attempts, NOW)
    assert sig.level == "weak" and sig.evidence_count == 3 and "SQL joins" in sig.explanation
    assert assessment_signal("Kubernetes", attempts, NOW) is UNASSESSED
